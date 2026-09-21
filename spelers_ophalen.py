"""Haalt spelers per team, bondsnummers en moyennes op van biljartpoint.nl
en schrijft ze naar Spelers_biljartpoint.xlsx. Dat bestand importeer je in de app via
Spelers -> "Importeer .xlsx / .csv".

De teamindeling komt uit de gespeelde partijen: biljartpoint geeft bij "teamleden" de hele
clubselectie per klasse (bijv. De Schalm 1 en 2 hebben dezelfde lijst), niet het echte team.
  - vaste speler van een team = speelde daar de meeste van zijn partijen
  - reserve                   = speelde ook (minder vaak) voor dat team
  - nog niet gespeeld          = staat onder "<club> (geen team)", alleen vindbaar als invaller

Schrijft ook spelers.json (voor de online app). Dat bestand wordt alleen herschreven als er echt
iets veranderd is (moyenne, team, nieuwe speler), zodat de app dan een "Bijwerken"-melding toont.

Gebruik:  python spelers_ophalen.py                    (standaard: competities 9830 en 9831)
          python spelers_ophalen.py 9830 9831 9900
          python spelers_ophalen.py --alleen-json      (zonder Excel-bestand, zo draait het op GitHub)
Nodig:    pip install openpyxl   (niet nodig met --alleen-json)
"""
import datetime
import hashlib
import json
import sys
import time
import urllib.request
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

API = 'https://biljartpoint.nl/api'
ALLEEN_JSON = '--alleen-json' in sys.argv
COMPETITIES = [int(a) for a in sys.argv[1:] if a.isdigit()] or [9830, 9831]
UIT = Path(__file__).with_name('Spelers_biljartpoint.xlsx')
JSON_UIT = Path(__file__).with_name('spelers.json')


def get(path):
    req = urllib.request.Request(API + path, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.load(r)


def te_maken(moy, factor, minimum, maximum):
    # Formule van de competitie: moyenne x factor, gewoon afgerond (0,5 naar boven), begrensd
    return max(minimum, min(maximum, int((Decimal(str(moy)) * factor).quantize(Decimal('1'), ROUND_HALF_UP))))


def naam_van(sp):
    return ' '.join(x for x in (sp.get('voornaam'), sp.get('tussenvoegsel'), sp.get('achternaam')) if x)


def main():
    # Alles over alle competities samen verzamelen, zodat een speler die in Klasse A speelt
    # niet in Klasse B als "geen team" of als vaste speler na één invalbeurt terechtkomt.
    spelers = {}                               # bondsnr -> naam, moy, club, formule
    gespeeld = defaultdict(Counter)            # bondsnr -> Counter(team -> aantal partijen)
    comp_van_team = {}
    for cid in COMPETITIES:
        data = get(f'/competities/{cid}')
        comp = data['competitie']
        lo, hi = comp.get('formule_min_caramboles') or 15, comp.get('formule_max_caramboles') or 300
        formule = comp.get('formule') or 'X*25'
        f = (int(formule[2:]) if formule.startswith('X*') else 25, lo, hi)

        # 1. selectie per team (moyenne, naam, club)
        for s in data['stand']:
            team = get(f"/teams/{s['team_id']}")
            time.sleep(0.1)
            comp_van_team[team['naam']] = comp['naam']
            for lid in team['team_leden']:
                bnr = int(lid['speler']['bondsnummer'])
                spelers.setdefault(bnr, {'naam': naam_van(lid['speler']), 'moy': lid['moyenne'],
                                         'club': team['vereniging']['naam'], 'f': f, 'comp': comp['naam']})

        # 2. wie speelde echt voor welk team (uit de partijen van gespeelde wedstrijden)
        klaar = [w for w in data['wedstrijden'] if w['status'] == 'afgerond' and not w.get('vervallen')]
        print(f"{cid}: {comp['naam']} – {len(data['stand'])} teams, {len(klaar)} gespeelde wedstrijden ophalen…")
        for w in klaar:
            for p in get(f"/wedstrijden/{w['id']}/partijen"):
                for kant in ('thuis', 'uit'):
                    sp = p.get(kant) or {}
                    if not sp.get('bondsnummer'):
                        continue
                    bnr = int(sp['bondsnummer'])
                    gespeeld[bnr][w[kant]['team_naam']] += 1
                    if bnr not in spelers:              # invaller die in geen enkele selectie staat
                        spelers[bnr] = {'naam': naam_van(sp.get('speler') or {}) or sp.get('speler_naam', ''),
                                        'moy': float(sp.get('huidige_moyenne') or 0), 'club': '', 'f': f, 'comp': comp['naam']}
            time.sleep(0.05)

    # 3. rijen: vaste speler bij het team waar hij de meeste partijen speelde, reserve bij de andere
    rijen = []
    for bnr, info in spelers.items():
        moy = info['moy'] or 0
        tm = te_maken(moy, *info['f'])
        teams = gespeeld.get(bnr)
        if teams:
            hoofd = teams.most_common(1)[0][0]
            for team, n in teams.items():
                rijen.append((comp_van_team.get(team, info['comp']), team, bnr, info['naam'], tm, moy, team != hoofd, n))
        else:
            rijen.append((info['comp'], f"{info['club']} (geen team)", bnr, info['naam'], tm, moy, True, 0))

    rijen.sort(key=lambda r: (r[1], r[6], -r[7], r[3]))
    schrijf_json(rijen)
    if not ALLEEN_JSON:
        schrijf_xlsx(rijen)


def schrijf_json(rijen):
    spelers = [[team, bnr, naam, tm, moy, 1 if res else 0] for _, team, bnr, naam, tm, moy, res, _ in rijen]
    # versie over een gesorteerde lijst: alleen echte wijzigingen tellen, niet een andere volgorde
    versie = hashlib.sha1(json.dumps(sorted(spelers, key=lambda r: (r[0], r[1])), ensure_ascii=False).encode('utf-8')).hexdigest()[:12]
    oud = json.loads(JSON_UIT.read_text(encoding='utf-8')) if JSON_UIT.exists() else {}
    if oud.get('version') == versie:
        print(f'{JSON_UIT.name}: geen wijzigingen (versie {versie})')
        return
    data = {'version': versie, 'bijgewerkt': datetime.date.today().isoformat(),
            'bron': 'biljartpoint.nl competities ' + ', '.join(map(str, COMPETITIES)), 'spelers': spelers}
    JSON_UIT.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'{JSON_UIT.name}: bijgewerkt naar versie {versie} ({len(spelers)} regels)')


def schrijf_xlsx(rijen):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    ws = wb.active
    ws.title = 'tblMoyenne'
    ws.append(['Teamnaam', 'Bondsnr', 'Naam', 'Temaken', 'Moy', 'Competitie', 'Reserve', 'Partijen'])
    for comp, team, bnr, naam, tm, moy, res, n in rijen:
        ws.append([team, bnr, naam, tm, moy, comp, 'ja' if res else '', n])
        ws.cell(ws.max_row, 5).number_format = '0.000'
    tabel = Table(displayName='tblMoyenne', ref=f'A1:H{ws.max_row}')
    tabel.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tabel)
    for col, w in zip('ABCDEFGH', [26, 10, 26, 10, 8, 12, 9, 9]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    wb.save(UIT)
    teams = {r[1] for r in rijen if not r[1].endswith('(geen team)')}
    print(f'{len(rijen)} regels, {len(teams)} teams opgeslagen in {UIT.name}')


if __name__ == '__main__':
    main()
